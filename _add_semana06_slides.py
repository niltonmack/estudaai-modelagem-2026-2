# -*- coding: utf-8 -*-
from copy import deepcopy
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SRC = r"c:\APLICACOES\CODIGO\estudaai-modelagem\Modelagem-Software-Semana06-SDD-v2.pptx"
DST = r"c:\APLICACOES\CODIGO\estudaai-modelagem\Modelagem-Software-Semana06-SDD-v3.pptx"

RED = RGBColor(0xD1, 0x50, 0x5B)
DARK = RGBColor(0x1F, 0x29, 0x33)
BODY = RGBColor(0x2B, 0x34, 0x40)
MUTED = RGBColor(0x5B, 0x66, 0x73)
HEAD_BG = RGBColor(0xEE, 0xF2, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF2, 0xF4, 0xF7)
LINE = RGBColor(0xFF, 0xFF, 0xFF)

SZ_KICKER = Emu(82550)
SZ_TITLE = Emu(355600)
SZ_LEAD = Emu(177800)
SZ_BODY = Emu(177800)
SZ_SMALL = Emu(152400)
SZ_NUM = Emu(82550)


def _run(p, text, *, size, bold=False, color=DARK, name="Arial"):
    r = p.add_run()
    r.text = text
    r.font.name = name
    r.font.size = size
    r.font.bold = bold
    r.font.color.rgb = color
    return r


def add_box(slide, l, t, w, h, fill=None, background=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.line.color.rgb = LINE
    sh.line.width = Emu(12700)
    if background:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill or WHITE
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    sh.text_frame.margin_left = Emu(91440)
    sh.text_frame.margin_right = Emu(91440)
    sh.text_frame.margin_top = Emu(68580)
    sh.text_frame.margin_bottom = Emu(45720)
    return sh


def set_anchor(sh, anchor=MSO_ANCHOR.MIDDLE):
    sh.text_frame._txBody.bodyPr.set("anchor", {MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.TOP: "t"}.get(anchor, "t"))


def clear_p(tf):
    p = tf.paragraphs[0]
    p.clear()
    return p


def header_block(slide, number, kicker, title, lead=None):
    num = add_box(slide, 11688775, 6556248, 228600, 146304, background=True)
    p = clear_p(num.text_frame)
    p.alignment = PP_ALIGN.CENTER
    _run(p, str(number), size=SZ_NUM, color=RGBColor(0x11, 0x11, 0x11))
    k = add_box(slide, 640080, 484632, 9000000, 164592, background=True)
    p = clear_p(k.text_frame)
    _run(p, kicker, size=SZ_KICKER, bold=True, color=RED)
    tit = add_box(slide, 640080, 804672, 10800000, 380000, background=True)
    p = clear_p(tit.text_frame)
    _run(p, title, size=SZ_TITLE, bold=True, color=DARK, name="Georgia")
    if lead:
        ld = add_box(slide, 640080, 1180000, 10800000, 320000, background=True)
        p = clear_p(ld.text_frame)
        _run(p, lead, size=SZ_LEAD, color=MUTED)
        return 1550000
    return 1250000


def footer_note(slide, text, top=5257800):
    sh = add_box(slide, 640080, top, 10800000, 420000, background=True)
    p = clear_p(sh.text_frame)
    _run(p, "• ", size=SZ_BODY, bold=True, color=RED)
    _run(p, text, size=SZ_BODY, color=BODY)


def table(slide, origin_y, col_widths, rows, header=True, row_h=520000, left=502920):
    y = origin_y
    for ri, row in enumerate(rows):
        x = left
        bg = HEAD_BG if (header and ri == 0) else WHITE
        for ci, cell in enumerate(row):
            w = col_widths[ci]
            sh = add_box(slide, x, y, w, row_h, fill=bg)
            set_anchor(sh, MSO_ANCHOR.MIDDLE)
            tf = sh.text_frame
            p = clear_p(tf)
            is_head = header and ri == 0
            is_first = ci == 0 and not is_head
            _run(
                p,
                cell,
                size=SZ_SMALL if len(cell) > 42 else SZ_BODY,
                bold=is_head or is_first,
                color=DARK if (is_head or is_first) else BODY,
            )
            x += w
        y += row_h
    return y


def cards_row(slide, y, items, h=1500000, left=685800, gap=137160, width=10607040):
    n = len(items)
    w = (width - gap * (n - 1)) // n
    x = left
    for title, body in items:
        sh = add_box(slide, x, y, w, h, fill=CARD)
        tf = sh.text_frame
        p = clear_p(tf)
        _run(p, title + "\n", size=SZ_BODY, bold=True, color=DARK)
        _run(p, body, size=SZ_SMALL, color=BODY)
        x += w + gap


def add_slide_at(prs, index):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    sldIdLst = prs.slides._sldIdLst
    new = list(sldIdLst)[-1]
    sldIdLst.remove(new)
    sldIdLst.insert(index, new)
    return prs.slides[index]


def renumber(prs):
    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if t.isdigit() and sh.left > 11000000:
                p = sh.text_frame.paragraphs[0]
                if p.runs:
                    p.runs[0].text = str(i)
                else:
                    p.clear()
                    _run(p, str(i), size=SZ_NUM, color=RGBColor(0x11, 0x11, 0x11))


def main():
    prs = Presentation(SRC)
    already = any(
        (sh.has_text_frame and sh.text_frame.text.strip() == "O que é MCP")
        for s in prs.slides
        for sh in s.shapes
    )
    if already:
        print("slides MCP já existem em", SRC, "- nada a inserir")
        return
    # Inserir depois de "Defesa em camadas" (índice 18) e antes da síntese (19).
    i = 19

    # --- MCP definição ---
    s = add_slide_at(prs, i)
    y = header_block(
        s,
        20,
        "FERRAMENTAS DO AGENTE",
        "O que é MCP",
        "MCP (Model Context Protocol) é um padrão aberto para o agente usar ferramentas externas — banco, GitHub, S3, SSH — sem colocar senha no prompt.",
    )
    cards_row(
        s,
        1600000,
        [
            (
                "O que conecta",
                "O modelo de linguagem não acessa MySQL ou S3 sozinho. Um servidor MCP expõe ferramentas que o harness pode chamar.",
            ),
            (
                "O que não é",
                "Não é a Spec, não é rule e não é o harness. É só o plugue de ferramentas externas.",
            ),
            (
                "Como o aluno usa",
                "Arquivos no repositório descrevem os servidores. A IDE sobe os processos ao abrir o projeto.",
            ),
        ],
        h=2000000,
    )
    footer_note(s, "Sem MCP, o agente só vê o que está no chat e nos arquivos do workspace.", top=5000000)
    i += 1

    # --- json vs env ---
    s = add_slide_at(prs, i)
    y = header_block(
        s,
        21,
        "DOIS ARQUIVOS",
        "mcp.json descreve ferramentas; mcp.env guarda valores",
    )
    table(
        s,
        y,
        [2600000, 5200000, 3000000],
        [
            ["Arquivo", "Função", "Vai para o Git?"],
            ["mcp.json", "Quais servidores existem, comando e nomes das variáveis", "Sim — sem senha"],
            ["mcp.env", "Host, usuário, senha, chave (placeholders no modelo)", "Só com sua-senha / seu-host"],
        ],
        row_h=700000,
        left=640080,
    )
    footer_note(
        s,
        "Na aula o modelo usa sua-senha, seu-host e sua-chave. Nunca coloque senha real no json.",
        top=4000000,
    )
    i += 1

    # --- 3 IDEs MCP ---
    s = add_slide_at(prs, i)
    y = header_block(s, 22, "TRÊS IDEs", "Como ligar mcp.json e mcp.env")
    table(
        s,
        y,
        [2400000, 4000000, 4200000],
        [
            ["IDE", "Arquivos", "Como json e env se ligam"],
            [
                "Cursor",
                ".cursor/mcp.json e .cursor/mcp.env",
                "envFile + ${env:MYSQL_HOST}",
            ],
            [
                "Claude Code",
                ".mcp.json na raiz e .claude/mcp.env",
                "${MYSQL_HOST} no ambiente (não lê o .env sozinho)",
            ],
            [
                "VS Code + Copilot",
                ".vscode/mcp.json (chave servers) e .vscode/mcp.env",
                "envFile + ${env:MYSQL_HOST}",
            ],
        ],
        row_h=780000,
        left=502920,
    )
    footer_note(s, "O json lista ferramentas; o env preenche valores. Os nomes das variáveis têm de coincidir.", top=5000000)
    i += 1

    # --- Rule vs Skill ---
    s = add_slide_at(prs, i)
    y = header_block(
        s,
        23,
        "INSTRUÇÃO DO AGENTE",
        "Rule é regulamento; skill é roteiro",
        "Rule: “neste projeto, faça assim”. Skill: “quando for fazer X, siga estes passos”.",
    )
    table(
        s,
        y,
        [2000000, 4600000, 4000000],
        [
            ["Critério", "Rule", "Skill"],
            ["O que é", "Convenção permanente do projeto", "Procedimento de uma tarefa"],
            ["Quando entra", "Sempre, por pasta (glob/paths) ou por descrição", "Sob demanda, ou chamada pelo nome"],
            ["Tamanho", "Curta: padrões e restrições", "Pode ser longa: fluxo, exemplos, scripts"],
            ["Analogia", "Regulamento da disciplina", "Roteiro de um laboratório"],
        ],
        row_h=580000,
        left=502920,
    )
    i += 1

    # --- Cursor vs Claude ---
    s = add_slide_at(prs, i)
    y = header_block(s, 24, "CURSOR E CLAUDE CODE", "Os dois papéis existem; as pastas mudam")
    table(
        s,
        y,
        [2800000, 4000000, 3800000],
        [
            ["Papel", "Cursor", "Claude Code"],
            ["Manual sempre ligado", "Rule alwaysApply ou AGENTS.md", ".claude/CLAUDE.md"],
            ["Rule (convenção)", ".cursor/rules/*.mdc", ".claude/rules/*.md (paths:)"],
            ["Skill (roteiro)", ".cursor/skills/<nome>/SKILL.md", ".claude/skills/<nome>/SKILL.md"],
            ["MCP", ".cursor/mcp.json", ".mcp.json na raiz do projeto"],
        ],
        row_h=680000,
        left=502920,
    )
    footer_note(
        s,
        "No Cursor, os 7 papéis estão em rules com alwaysApply: false. No Claude Code o equivalente é skill.",
        top=5000000,
    )
    i += 1

    # --- VS Code ---
    s = add_slide_at(prs, i)
    y = header_block(
        s,
        25,
        "VS CODE + COPILOT",
        "Funciona de outro jeito — use o modelo no GitHub",
        "O Copilot não cabe só em .vscode. MCP fica ali; o manual e as rules ficam em .github/.",
    )
    table(
        s,
        y,
        [3600000, 7000000],
        [
            ["Papel no Copilot", "Onde configurar no EstudaAI"],
            ["Manual sempre ligado", ".github/copilot-instructions.md"],
            ["Rule por pasta", ".github/instructions/*.instructions.md (applyTo)"],
            ["Skills do pipeline", "O Copilot já lê .claude/skills/ (não duplicar)"],
            ["MCP", ".vscode/mcp.json (chave servers) + .vscode/mcp.env"],
        ],
        row_h=560000,
        left=640080,
    )
    footer_note(
        s,
        "Baixe o modelo no repositório GitHub do EstudaAI: pastas .vscode, .github, .cursor e .claude. Comece por .vscode/COPILOT.md.",
        top=5000000,
    )
    i += 1

    # --- Harness ---
    s = add_slide_at(prs, i)
    y = header_block(
        s,
        26,
        "RUNTIME DO AGENTE",
        "O harness é quem lê esses arquivos",
        "Harness (ambiente de execução) envolve o modelo. Não é o LLM. É o Cursor, o Claude Code ou o Agent Host do VS Code/Copilot.",
    )
    cards_row(
        s,
        y,
        [
            (
                "Carrega contexto",
                "Injeta CLAUDE.md, copilot-instructions e rules no início da sessão — ou quando a pasta casa com glob/paths.",
            ),
            (
                "Descobre skills",
                "Só o índice (name + description) fica sempre visível. O corpo do SKILL.md entra quando a tarefa combina.",
            ),
            (
                "Sobe o MCP",
                "Lê mcp.json, interpola variáveis, inicia os servidores e decide se pede confiança ao aluno.",
            ),
            (
                "Intercepta ferramentas",
                "Hooks, confirmação e trust atuam fora do modelo. Por isso rule não substitui gate de segurança.",
            ),
        ],
        h=2200000,
        width=10800000,
        left=640080,
        gap=100000,
    )
    footer_note(
        s,
        "Harnesses diferentes leem pastas diferentes. A Spec continua mandando no comportamento; o harness só decide o que o modelo enxerga e o que pode executar.",
        top=5000000,
    )

    # Síntese (índice 26): amarrar MCP, rules/skills e harness ao recap.
    sint = prs.slides[26]
    for sh in sint.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("O agente opera com instruções"):
            p = sh.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = (
                    "Rules/skills e MCP entram via harness; hooks, Git e CI continuam fora do modelo"
                )

    renumber(prs)
    prs.save(DST)
    print("saved", DST, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
